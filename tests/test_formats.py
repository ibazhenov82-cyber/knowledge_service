"""Извлечение текста и структуры из всех поддерживаемых форматов."""

from __future__ import annotations

import io
import json
import os
import unittest

from knowledge_service.code_structure import split_code
from knowledge_service.formats import (
    ExtractError, FileTypes, parse_document, resolve_kind,
)

FT = FileTypes()


def parse(name: str, data, file_types: FileTypes = FT):
    if isinstance(data, str):
        data = data.encode("utf-8")
    return parse_document(data, name, file_types)


MD = """---
title: Регламент ревью
date: 2026-05-12
version: 2.1
author: Иван
tags: [ревью, процесс]
---
# Регламент

Вступление.

## Сроки ревью

Ревью выполняется не позднее 2 рабочих дней.

```python
# Это не заголовок
x = 1
```

### Исключения

Срочные исправления — в тот же день.

#### Глубокий заголовок остаётся внутри

Текст.

Второй раздел
-------------

Текст второго раздела.
"""


class TextFormatsTests(unittest.TestCase):
    def test_markdown_sections_and_front_matter(self):
        p = parse("review.md", MD)
        self.assertEqual(p.title, "Регламент ревью")
        self.assertEqual(p.doc_date, "2026-05-12")
        self.assertEqual(p.doc_version, "2.1")
        self.assertEqual(p.properties, {"author": "Иван", "tags": "ревью, процесс", "sections": 4})
        sections = [u.section for u in p.units]
        self.assertEqual(sections, ["Регламент", "Регламент › Сроки ревью",
                                    "Регламент › Сроки ревью › Исключения", "Регламент › Второй раздел"])
        self.assertIn("# Это не заголовок", p.units[1].text)
        self.assertIn("Глубокий заголовок", p.units[2].text)

    def test_rst_headings(self):
        text = "=====\nГлава\n=====\n\nВведение.\n\nРаздел 1\n--------\n\nТекст 1.\n\nПодраздел\n~~~~~~~~~\n\nТекст.\n"
        p = parse("doc.rst", text)
        self.assertEqual(p.title, "Глава")
        self.assertEqual([u.section for u in p.units], ["Глава", "Глава › Раздел 1", "Глава › Раздел 1 › Подраздел"])

    def test_html_main_content_and_skips(self):
        html = """<html lang="ru"><head><title>Wiki: Ревью</title><style>.x{}</style>
        <meta name="description" content="Как проводим ревью">
        <meta property="article:modified_time" content="2026-04-01T10:00:00Z"></head><body>
        <nav>Меню Меню Меню</nav>
        <main><h1>Ревью</h1><p>Общие правила ревью кода в команде, обязательные для всех репозиториев.</p>
        <h2>Сроки</h2><p>Не позднее <b>2</b> рабочих дней с момента запроса ревью, включая исправления.</p>
        <pre>git push
  --force</pre>
        <table><tr><th>Тип</th><th>Срок</th></tr><tr><td>Срочно</td><td>1 день</td></tr></table>
        </main><footer>Подвал</footer><script>alert(1)</script></body></html>"""
        p = parse("page.html", html)
        self.assertEqual(p.title, "Wiki: Ревью")
        self.assertEqual(p.properties["description"], "Как проводим ревью")
        self.assertEqual(p.properties["lang"], "ru")
        self.assertEqual(p.doc_date, "2026-04-01")
        joined = "\n".join(u.text for u in p.units)
        self.assertNotIn("Меню", joined)
        self.assertNotIn("Подвал", joined)
        self.assertNotIn("alert", joined)
        self.assertEqual([u.section for u in p.units], ["Ревью", "Ревью › Сроки"])
        self.assertIn("Не позднее 2 рабочих дней", joined)
        self.assertIn("  --force", joined)
        self.assertIn("Срочно | 1 день", joined)

    def test_text_and_encodings(self):
        p = parse_document("Привет, мир".encode("cp1251"), "note.txt", FT)
        self.assertEqual(p.units[0].text, "Привет, мир")
        self.assertEqual(p.title, "note")

    def test_extra_text_extension_and_unsupported(self):
        ft = FileTypes.from_dict({"extra_text_extensions": ["adoc", ".LOG"]})
        self.assertEqual(ft.extra_text_extensions, [".adoc", ".log"])
        self.assertEqual(resolve_kind("a.log", ft), "text")
        self.assertIsNone(resolve_kind("a.log", FT))
        with self.assertRaises(ExtractError) as ctx:
            parse("image.png", b"\x89PNG")
        self.assertIn("не поддерживается", str(ctx.exception))

    def test_configurable_code_extensions(self):
        ft = FileTypes.from_dict({"code_extensions": [".py", ".kt"]})
        self.assertEqual(resolve_kind("a.kt", ft), "code")
        self.assertIsNone(resolve_kind("a.go", ft))
        with self.assertRaises(ValueError):
            FileTypes.from_dict({"code_extensions": ["../x"]})

    def test_empty_document(self):
        with self.assertRaises(ExtractError):
            parse("empty.md", "   \n")


class JsonlTests(unittest.TestCase):
    DATA = "\n".join([
        json.dumps({"id": "T-1", "subject": "Не работает вход", "description": "Ошибка 500", "status": "closed",
                    "created_at": "2026-01-02"}, ensure_ascii=False),
        "не json",
        json.dumps({"id": "T-2", "subject": "Медленно", "description": "Долго грузится", "status": "open"},
                   ensure_ascii=False),
    ])

    def test_default_fields(self):
        p = parse("tickets.jsonl", self.DATA)
        self.assertEqual(len(p.units), 2)
        self.assertTrue(all(u.atomic for u in p.units))
        self.assertEqual(p.units[0].section, "Запись T-1")
        self.assertEqual(p.properties["records"], 2)
        self.assertEqual(p.units[0].metadata, {"id": "T-1"})
        self.assertIn("subject: Не работает вход", p.units[0].text)
        self.assertEqual(p.warnings, ["пропущено строк с некорректным JSON: 1"])

    def test_configured_fields(self):
        ft = FT.with_jsonl({"text_fields": ["subject", "description"], "metadata_fields": ["status", "created_at"]})
        p = parse("tickets.jsonl", self.DATA, ft)
        self.assertEqual(p.units[0].metadata, {"status": "closed", "created_at": "2026-01-02", "id": "T-1"})
        self.assertNotIn("status", p.units[0].text)


class OfficeFormatsTests(unittest.TestCase):
    def test_docx(self):
        import docx

        d = docx.Document()
        d.core_properties.title = "Политика отпусков"
        d.core_properties.author = "Отдел кадров"
        d.core_properties.keywords = "отпуск, кадры"
        d.core_properties.version = "3"
        d.add_heading("Отпуска", level=1)
        d.add_paragraph("Ежегодный отпуск — 28 дней.")
        d.add_heading("Перенос", level=2)
        d.add_paragraph("Перенос согласуется с руководителем.")
        table = d.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "Стаж", "Дни"
        table.cell(1, 0).text, table.cell(1, 1).text = "5 лет", "+3"
        buf = io.BytesIO()
        d.save(buf)
        p = parse("policy.docx", buf.getvalue())
        self.assertEqual(p.title, "Политика отпусков")
        self.assertEqual([u.section for u in p.units], ["Отпуска", "Отпуска › Перенос"])
        self.assertIn("5 лет | +3", p.units[1].text)
        self.assertIsNotNone(p.doc_date)
        self.assertEqual(p.doc_version, "3")
        self.assertEqual((p.properties["author"], p.properties["keywords"], p.properties["tables"]),
                         ("Отдел кадров", "отпуск, кадры", 1))

    def test_pptx(self):
        from pptx import Presentation

        prs = Presentation()
        for title, body in (("Цели", "Рост на 20%"), ("План", "Три этапа")):
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text = body
        prs.slides[1].notes_slide.notes_text_frame.text = "Сказать про сроки"
        buf = io.BytesIO()
        prs.save(buf)
        p = parse("deck.pptx", buf.getvalue())
        self.assertEqual([u.section for u in p.units], ["Слайд 1: Цели", "Слайд 2: План"])
        self.assertEqual([u.page for u in p.units], [1, 2])
        self.assertIn("Заметки: Сказать про сроки", p.units[1].text)
        self.assertEqual(p.properties["slides"], 2)

    def test_xlsx(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "Тарифы"
        ws.append(["Тариф", "Цена"])
        ws.append(["Базовый", 100])
        ws.append(["Про", 250.0])
        wb.create_sheet("Пусто")
        buf = io.BytesIO()
        wb.save(buf)
        p = parse("prices.xlsx", buf.getvalue())
        self.assertEqual(len(p.units), 1)
        self.assertEqual(p.units[0].section, "Лист «Тарифы»")
        self.assertIn("Тариф: Про; Цена: 250", p.units[0].text)
        self.assertEqual((p.properties["sheets"], p.properties["sheet_names"], p.properties["rows"]),
                         (2, "Тарифы, Пусто", 2))

    def test_broken_office_file(self):
        with self.assertRaises(ExtractError):
            parse("bad.docx", b"not a zip")


try:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    FONT = next((p for p in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                             "/usr/share/fonts/dejavu/DejaVuSans.ttf") if os.path.exists(p)), None)
    REPORTLAB = FONT is not None
except ImportError:  # pragma: no cover
    REPORTLAB = False


@unittest.skipUnless(REPORTLAB, "нет reportlab или шрифта DejaVu")
class PdfTests(unittest.TestCase):
    def make_pdf(self, pages):
        pdfmetrics.registerFont(TTFont("DejaVu", FONT))
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        c.setAuthor("Отдел качества")
        c.setSubject("Ревью кода")
        for lines in pages:
            c.setFont("DejaVu", 11)
            y = 800
            c.drawString(50, 820, "ООО Ромашка — внутренний документ")
            for line in lines:
                c.drawString(50, y, line)
                y -= 16
            c.drawString(50, 30, f"Страница {pages.index(lines) + 1}")
            c.showPage()
        c.save()
        return buf.getvalue()

    def test_pages_headings_and_repeated_headers(self):
        data = self.make_pdf([
            ["1 Общие положения", "Документ описывает порядок ревью.", "Ревью обязательно для всех."],
            ["2 Сроки", "Не позднее двух рабочих дней."],
            ["2.1 Исключения", "Срочные исправления — в тот же день."],
        ])
        p = parse("rules.pdf", data)
        self.assertEqual([u.page for u in p.units], [1, 2, 3])
        self.assertEqual(p.units[1].section, "2 Сроки")
        self.assertEqual(p.units[2].section, "2 Сроки › 2.1 Исключения")
        joined = "\n".join(u.text for u in p.units)
        self.assertNotIn("Ромашка", joined)
        self.assertNotIn("Страница", joined)
        self.assertEqual((p.properties["pages"], p.properties["author"], p.properties["subject"]),
                         (3, "Отдел качества", "Ревью кода"))
        self.assertIn("created", p.properties)

    def test_scan_without_text(self):
        buf = io.BytesIO()
        c = canvas.Canvas(buf)
        c.rect(10, 10, 100, 100)
        c.showPage()
        c.save()
        with self.assertRaises(ExtractError) as ctx:
            parse("scan.pdf", buf.getvalue())
        self.assertIn("OCR", str(ctx.exception))


class CodeTests(unittest.TestCase):
    def sections(self, name, src):
        language, units = split_code(src, name)
        return language, [u.section for u in units], units

    def test_python(self):
        src = '''"""Модуль."""
import os

# Константа
LIMIT = 3


@decorator
def top(a):
    return a


class Repo:
    """Репозиторий."""

    x = 1

    def load(self):
        return 1

    # Сохранить
    async def save(self):
        return 2


if __name__ == "__main__":
    top(1)
'''
        language, sections, units = self.sections("repo.py", src)
        self.assertEqual(language, "python")
        self.assertEqual(sections, ["", "top", "Repo", "Repo › load", "Repo › save", ""])
        self.assertTrue(units[1].text.startswith("@decorator"))
        self.assertTrue(units[4].text.startswith("    # Сохранить"))

    def test_python_syntax_error_falls_back(self):
        _, sections, units = self.sections("bad.py", "def (:\n  pass\n")
        self.assertEqual(sections, [""])

    def test_kotlin(self):
        src = """package app

import foo.Bar

/** Модель. */
data class User(val id: String)

class Repo(private val api: Api) {
    fun load(): User = api.get()

    suspend fun save(user: User) {
        api.put(user)
    }

    companion object {
        const val TAG = "Repo"
    }
}

fun main() {
    println("x")
}
"""
        language, sections, _ = self.sections("Repo.kt", src)
        self.assertEqual(language, "kotlin")
        self.assertEqual(sections, ["", "User", "Repo", "Repo › load", "Repo › save", "Repo › companion", "main"])

    def test_java(self):
        src = """package app;

public class Service {
    private final Repo repo;

    @Override
    public String toString() {
        return "s";
    }

    public List<User> findAll(int limit) {
        if (limit > 0) {
            return repo.all();
        }
        return List.of();
    }
}
"""
        _, sections, units = self.sections("Service.java", src)
        self.assertEqual(sections, ["", "Service", "Service › toString", "Service › findAll"])
        self.assertTrue(units[2].text.strip().startswith("@Override"))

    def test_typescript(self):
        src = """import { x } from "./x";

export interface Props {
  id: string;
}

export class Store {
  private items: string[] = [];

  async load(id: string): Promise<void> {
    if (id) {
      return;
    }
  }
}

export const helper = (a: number) => a + 1;

export default function App() {
  return null;
}
"""
        _, sections, _ = self.sections("store.ts", src)
        self.assertEqual(sections, ["", "Props", "Store", "Store › load", "helper", "App"])

    def test_go(self):
        src = "package main\n\nimport \"fmt\"\n\n// User — пользователь.\ntype User struct{}\n\n" \
              "func (u *User) Name() string { return \"\" }\n\nfunc main() {\n\tfmt.Println(1)\n}\n"
        _, sections, units = self.sections("main.go", src)
        self.assertEqual(sections, ["", "User", "User.Name", "main"])
        self.assertTrue(units[1].text.startswith("// User"))

    def test_sql_yaml_json_toml(self):
        _, sections, _ = self.sections("schema.sql", "-- Пользователи\nCREATE TABLE users (\n  id int\n);\n\n"
                                                      "CREATE INDEX users_id ON users(id);\nSELECT * FROM users;\n")
        self.assertEqual(sections, ["CREATE TABLE USERS", "CREATE INDEX USERS_ID", "SELECT *"])
        _, sections, _ = self.sections("ci.yaml", "# CI\nname: build\non:\n  push: {}\njobs:\n  test:\n    runs: x\n")
        self.assertEqual(sections, ["name", "on", "jobs"])
        _, sections, units = self.sections("cfg.json", json.dumps({"server": {"port": 1}, "db": "x"}))
        self.assertEqual(sections, ["server", "db"])
        self.assertIn('"port": 1', units[0].text)
        _, sections, _ = self.sections("pyproject.toml", "name = 'x'\n\n[tool.black]\nline = 1\n\n[[bin]]\nname='a'\n")
        self.assertEqual(sections, ["", "tool.black", "bin"])

    def test_generic_extension_whole_file(self):
        _, sections, units = self.sections("run.sh", "#!/bin/sh\necho 1\n")
        self.assertEqual(sections, [""])
        p = parse("Repo.kt", "class A\n")
        self.assertEqual(p.doc_type, "code")
        self.assertEqual(p.code_language, "kotlin")
        self.assertEqual((p.properties["language"], p.properties["lines"], p.properties["declarations"]), ("kotlin", 2, 1))
        self.assertEqual(p.title, "Repo.kt")


if __name__ == "__main__":
    unittest.main()
