"""
knowledge_service.formats
===========================

Определение формата файла и извлечение текста со структурой.

Результат разбора — `Parsed`: заголовок документа и последовательность
структурных единиц (`Unit`): разделов по заголовкам, страниц, слайдов,
листов, объявлений кода, записей JSONL. Способ «По структуре» режет по этим
единицам, «Фиксированный размер» — по склеенному из них тексту, беря из
единиц раздел и страницу.
"""

from __future__ import annotations

import io
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from datetime import datetime
from html.parser import HTMLParser
from pathlib import PurePosixPath
from typing import Any, Dict, List, Optional, Tuple

from . import code_structure

TEXT_EXTENSIONS = [".txt", ".md", ".markdown", ".rst", ".html", ".htm"]
DOCUMENT_EXTENSIONS = [".pdf", ".docx", ".pptx", ".xlsx"]
STRUCTURED_EXTENSIONS = [".jsonl"]
DEFAULT_CODE_EXTENSIONS = [
    ".py", ".pyi", ".kt", ".kts", ".java", ".js", ".ts", ".tsx", ".jsx", ".go", ".sql",
    ".yaml", ".yml", ".json", ".toml", ".xml", ".gradle", ".properties", ".sh",
]

SECTION_SEP = " › "

#: Тип по Content-Type — для ссылок без расширения.
CONTENT_TYPES = {
    "text/html": ".html", "application/xhtml+xml": ".html", "text/plain": ".txt", "text/markdown": ".md",
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/json": ".json", "application/x-ndjson": ".jsonl", "application/jsonl": ".jsonl",
}


class ExtractError(Exception):
    """Документ не удалось разобрать (причина — для пользователя)."""


@dataclass
class Unit:
    text: str
    section: str = ""
    page: Optional[int] = None
    #: Не склеивать с соседними (запись JSONL).
    atomic: bool = False
    #: Метаданные, которые переходят во фрагменты (поля записи JSONL).
    metadata: Dict[str, Any] = field(default_factory=dict)


@dataclass
class Parsed:
    title: str
    units: List[Unit]
    doc_type: str
    doc_date: Optional[str] = None
    #: Язык кода для файлов кода.
    code_language: Optional[str] = None
    #: Версия документа из свойств файла (front matter, свойства DOCX).
    doc_version: Optional[str] = None
    warnings: List[str] = field(default_factory=list)
    #: Свойства, извлечённые из самого файла: автор, число страниц, тема,
    #: ключевые слова, поля front matter, <meta> страницы и т. п. Только
    #: простые значения (строки, числа).
    properties: Dict[str, Any] = field(default_factory=dict)


def _prop(value: Any) -> Any:
    """Значение свойства: строка/число; даты — ГГГГ-ММ-ДД, списки — через запятую."""
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, bool) or isinstance(value, (int, float)):
        return value
    if isinstance(value, (list, tuple)):
        value = ", ".join(str(v) for v in value if v is not None)
    text = str(value).strip()
    # Заглушки генераторов PDF/офисных файлов вместо пустого значения.
    if not text or text.lower() in _PLACEHOLDERS:
        return None
    return text[:500]


_PLACEHOLDERS = {"unspecified", "anonymous", "untitled", "unknown", "none", "null", "(anonymous)", "(unspecified)"}


def _props(**values: Any) -> Dict[str, Any]:
    result = {}
    for key, value in values.items():
        clean = _prop(value)
        if clean is not None and clean != "":
            result[key] = clean
    return result


# ---------------------------------------------------------------------------
# Настройки типов файлов коллекции
# ---------------------------------------------------------------------------


def _norm_ext(ext: str) -> str:
    ext = ext.strip().lower()
    if not ext:
        raise ValueError("пустое расширение")
    if not ext.startswith("."):
        ext = "." + ext
    if not re.fullmatch(r"\.[a-z0-9_+-]{1,20}", ext):
        raise ValueError(f"некорректное расширение {ext!r}")
    return ext


@dataclass
class FileTypes:
    code_extensions: List[str] = field(default_factory=lambda: list(DEFAULT_CODE_EXTENSIONS))
    extra_text_extensions: List[str] = field(default_factory=list)
    #: JSONL: какие поля записи — текст (пусто — все строковые поля, кроме
    #: метаданных) и какие — метаданные фрагмента.
    jsonl_text_fields: List[str] = field(default_factory=list)
    jsonl_metadata_fields: List[str] = field(default_factory=list)

    @classmethod
    def from_dict(cls, data: Optional[Dict[str, Any]], base: Optional["FileTypes"] = None) -> "FileTypes":
        base = base or cls()
        if not data:
            return cls(**base.to_dict_flat())
        if not isinstance(data, dict):
            raise ValueError("file_types должен быть объектом")
        unknown = set(data) - {"code_extensions", "extra_text_extensions", "jsonl"}
        if unknown:
            raise ValueError(f"file_types: неизвестные поля {sorted(unknown)}")
        result = cls(**base.to_dict_flat())
        if "code_extensions" in data:
            result.code_extensions = sorted({_norm_ext(e) for e in data["code_extensions"] or []})
        if "extra_text_extensions" in data:
            result.extra_text_extensions = sorted({_norm_ext(e) for e in data["extra_text_extensions"] or []})
        if "jsonl" in data:
            result.jsonl_text_fields, result.jsonl_metadata_fields = parse_jsonl_options(data["jsonl"])
        return result

    def to_dict_flat(self) -> Dict[str, Any]:
        return {
            "code_extensions": list(self.code_extensions), "extra_text_extensions": list(self.extra_text_extensions),
            "jsonl_text_fields": list(self.jsonl_text_fields), "jsonl_metadata_fields": list(self.jsonl_metadata_fields),
        }

    def to_dict(self) -> Dict[str, Any]:
        return {
            "code_extensions": list(self.code_extensions), "extra_text_extensions": list(self.extra_text_extensions),
            "jsonl": {"text_fields": list(self.jsonl_text_fields), "metadata_fields": list(self.jsonl_metadata_fields)},
        }

    def with_jsonl(self, options: Optional[Dict[str, Any]]) -> "FileTypes":
        if not options:
            return self
        result = FileTypes(**self.to_dict_flat())
        result.jsonl_text_fields, result.jsonl_metadata_fields = parse_jsonl_options(options)
        return result


def parse_jsonl_options(options: Any) -> Tuple[List[str], List[str]]:
    if not isinstance(options, dict) or set(options) - {"text_fields", "metadata_fields"}:
        raise ValueError("jsonl: ожидается {\"text_fields\": [...], \"metadata_fields\": [...]}")
    text_fields = [str(f).strip() for f in options.get("text_fields") or [] if str(f).strip()]
    meta_fields = [str(f).strip() for f in options.get("metadata_fields") or [] if str(f).strip()]
    return text_fields, meta_fields


def extension_of(filename: str) -> str:
    return PurePosixPath(filename.replace("\\", "/")).suffix.lower()


def resolve_kind(filename: str, file_types: FileTypes) -> Optional[str]:
    """Вид разбора по имени файла; None — формат не поддерживается."""
    ext = extension_of(filename)
    if ext in (".md", ".markdown"):
        return "markdown"
    if ext == ".rst":
        return "rst"
    if ext in (".html", ".htm"):
        return "html"
    if ext == ".txt":
        return "text"
    if ext in (".pdf", ".docx", ".pptx", ".xlsx"):
        return ext[1:]
    if ext == ".jsonl":
        return "jsonl"
    if ext in file_types.code_extensions:
        return "code"
    if ext in file_types.extra_text_extensions:
        return "text"
    return None


def supported_formats(file_types: Optional[FileTypes] = None) -> Dict[str, List[str]]:
    file_types = file_types or FileTypes()
    return {
        "text": TEXT_EXTENSIONS, "documents": DOCUMENT_EXTENSIONS,
        "code": list(file_types.code_extensions), "structured": STRUCTURED_EXTENSIONS,
        "extra_text": list(file_types.extra_text_extensions),
    }


# ---------------------------------------------------------------------------
# Текст
# ---------------------------------------------------------------------------


def decode_text(data: bytes, charset: Optional[str] = None) -> str:
    if data.startswith(b"\xef\xbb\xbf"):
        data = data[3:]
    for enc in [c for c in (charset, "utf-8") if c]:
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    try:
        return data.decode("cp1251")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _stem(filename: str) -> str:
    return PurePosixPath(filename.replace("\\", "/")).stem or filename


class _SectionPath:
    """Путь раздела по заголовкам уровней 1–3."""

    def __init__(self) -> None:
        self.levels: List[str] = []

    def set(self, level: int, title: str) -> str:
        level = max(1, min(level, 3))
        self.levels = self.levels[: level - 1] + [""] * max(0, level - 1 - len(self.levels))
        self.levels.append(title.strip())
        return self.path

    @property
    def path(self) -> str:
        return SECTION_SEP.join(t for t in self.levels if t)


_MD_HEADING_MARKS = re.compile(r"^\s*#{1,6}\s+|\s+#+\s*$")


def _units_from_headed_lines(items: List[Tuple[Optional[int], str]]) -> List[Unit]:
    """items: (уровень заголовка или None, строка). Новый раздел — на
    заголовках уровней 1–3; более глубокие остаются внутри раздела."""
    units: List[Unit] = []
    path = _SectionPath()
    section, buf = "", []

    def flush() -> None:
        text = _clean("\n".join(buf))
        if text:
            units.append(Unit(text=text, section=section))

    for level, line in items:
        if level is not None and level <= 3:
            flush()
            buf = []
            section = path.set(level, _MD_HEADING_MARKS.sub("", line).strip())
            buf.append(line)
        else:
            buf.append(line)
    flush()
    return units


_MD_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")


def parse_markdown(text: str, filename: str) -> Parsed:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    doc_date, fm_title = None, None
    front: Dict[str, Any] = {}
    if lines and lines[0].strip() == "---":
        for end in range(1, min(len(lines), 60)):
            if lines[end].strip() in ("---", "..."):
                for raw in lines[1:end]:
                    if raw.startswith((" ", "\t", "-")) or ":" not in raw:
                        continue
                    key, _, value = raw.partition(":")
                    key = key.strip()
                    value = value.strip().strip("'\"")
                    if not re.fullmatch(r"[\w.-]{1,40}", key) or not value:
                        continue
                    if key.lower() == "title":
                        fm_title = value
                    elif key.lower() in ("date", "updated", "last_modified"):
                        doc_date = value[:10]
                    else:
                        front[key] = value.strip("[]")
                lines = lines[end + 1:]
                break
    items: List[Tuple[Optional[int], str]] = []
    in_fence = False
    title = fm_title
    for i, line in enumerate(lines):
        if _MD_FENCE.match(line):
            in_fence = not in_fence
            items.append((None, line))
            continue
        if in_fence:
            items.append((None, line))
            continue
        match = _MD_HEADING.match(line)
        if match:
            level, heading = len(match.group(1)), match.group(2)
            title = title or (heading if level == 1 else None)
            items.append((level, line))
            continue
        # Setext-заголовки: строка текста, подчёркнутая === или ---.
        if i + 1 < len(lines) and line.strip() and not line.startswith((" ", "\t", "-", "*", ">", "|")):
            nxt = lines[i + 1].strip()
            if nxt and set(nxt) <= {"="} or (nxt and set(nxt) <= {"-"} and len(nxt) >= 3):
                level = 1 if nxt[0] == "=" else 2
                title = title or (line.strip() if level == 1 else None)
                items.append((level, line))
                continue
        if items and items[-1][0] is not None and line.strip() and set(line.strip()) <= {"=", "-"}:
            continue  # подчёркивание setext-заголовка
        items.append((None, line))
    version = front.pop("version", None) or front.pop("doc_version", None)
    return Parsed(title=title or _stem(filename), units=_units_from_headed_lines(items),
                  doc_type="markdown", doc_date=doc_date, doc_version=version, properties=_props(**front))


_RST_ADORN = set("=-`:'\"~^_*+#<>")


def parse_rst(text: str, filename: str) -> Parsed:
    lines = _clean(text).split("\n") if text.strip() else []
    styles: List[Tuple[str, bool]] = []
    items: List[Tuple[Optional[int], str]] = []
    title = None
    i = 0

    def is_adorn(line: str, min_len: int) -> bool:
        s = line.rstrip()
        return len(s) >= max(3, min_len) and len(set(s)) == 1 and s[0] in _RST_ADORN

    while i < len(lines):
        line = lines[i]
        over = False
        if is_adorn(line, 1) and i + 2 < len(lines) and lines[i + 1].strip() \
                and is_adorn(lines[i + 2], len(lines[i + 1].strip())) and lines[i + 2].strip()[0] == line.strip()[0]:
            over, heading, char = True, lines[i + 1].strip(), line.strip()[0]
            i += 3
        elif line.strip() and i + 1 < len(lines) and is_adorn(lines[i + 1], len(line.strip())) \
                and not line.startswith((" ", "\t")):
            heading, char = line.strip(), lines[i + 1].strip()[0]
            i += 2
        else:
            items.append((None, line))
            i += 1
            continue
        style = (char, over)
        if style not in styles:
            styles.append(style)
        level = styles.index(style) + 1
        title = title or (heading if level == 1 else None)
        items.append((level, heading))
    return Parsed(title=title or _stem(filename), units=_units_from_headed_lines(items), doc_type="rst")


def parse_text(text: str, filename: str, doc_type: str = "text") -> Parsed:
    text = _clean(text)
    return Parsed(title=_stem(filename), units=[Unit(text=text)] if text else [], doc_type=doc_type)


# ---------------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------------

_HTML_SKIP = {"script", "style", "noscript", "svg", "template", "iframe", "nav", "footer", "form", "button",
              "select", "aside"}
_HTML_BLOCK = {"p", "div", "br", "li", "ul", "ol", "table", "section", "article", "main", "blockquote",
               "header", "dl", "dt", "dd", "figure", "figcaption", "hr"}
#: <meta name=...> → свойство документа.
_HTML_META = {
    "description": "description", "og:description": "description", "author": "author",
    "keywords": "keywords", "article:published_time": "published", "article:modified_time": "modified",
    "og:site_name": "site",
}
_HTML_HEADINGS = {"h1": 1, "h2": 2, "h3": 3, "h4": 4, "h5": 5, "h6": 6}


class _HtmlBlocks(HTMLParser):
    """Последовательность блоков (уровень заголовка или None, текст) — отдельно
    для всей страницы и для <main>/<article>."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.skip = 0
        self.pre = 0
        self.main = 0
        self.title = ""
        self.in_title = False
        self.heading: Optional[int] = None
        self.buf: List[str] = []
        self.all: List[Tuple[Optional[int], str]] = []
        self.main_blocks: List[Tuple[Optional[int], str]] = []
        self.links: List[str] = []
        self.meta: Dict[str, str] = {}

    def _flush(self) -> None:
        raw = "".join(self.buf)
        self.buf = []
        if self.pre:
            text = raw.strip("\n")
        else:
            text = "\n".join(re.sub(r"[ \t\r\f\v ]+", " ", ln).strip() for ln in raw.split("\n"))
            text = re.sub(r"\n{2,}", "\n", text).strip()
        if not text:
            return
        block = (self.heading, text if self.heading is None else re.sub(r"\s+", " ", text))
        self.all.append(block)
        if self.main:
            self.main_blocks.append(block)

    def handle_starttag(self, tag, attrs):
        if tag in _HTML_SKIP:
            self.skip += 1
            return
        if self.skip:
            return
        if tag == "title":
            self.in_title = True
        elif tag == "a":
            href = dict(attrs).get("href")
            if href:
                self.links.append(href)
        elif tag == "meta":
            attr = dict(attrs)
            name = (attr.get("name") or attr.get("property") or "").lower()
            if name in _HTML_META and attr.get("content"):
                self.meta[_HTML_META[name]] = attr["content"].strip()
        elif tag == "html":
            lang = dict(attrs).get("lang")
            if lang:
                self.meta["lang"] = lang
        elif tag in ("main", "article"):
            self._flush()
            self.main += 1
        elif tag in _HTML_HEADINGS:
            self._flush()
            self.heading = _HTML_HEADINGS[tag]
        elif tag == "pre":
            self._flush()
            self.pre += 1
            self.buf.append("```\n")
        elif tag in ("td", "th"):
            self.buf.append(" | ")
        elif tag == "tr":
            self.buf.append("\n")
        elif tag in _HTML_BLOCK:
            if tag == "li":
                self.buf.append("\n- ")
            else:
                self._flush()

    def handle_endtag(self, tag):
        if tag in _HTML_SKIP:
            self.skip = max(0, self.skip - 1)
            return
        if self.skip:
            return
        if tag == "title":
            self.in_title = False
        elif tag in _HTML_HEADINGS:
            self._flush()
            self.heading = None
        elif tag == "pre":
            self.buf.append("\n```")
            self._flush()
            self.pre = max(0, self.pre - 1)
        elif tag in ("main", "article"):
            self._flush()
            self.main = max(0, self.main - 1)
        elif tag in _HTML_BLOCK and tag != "li":
            self._flush()

    def handle_data(self, data):
        if self.in_title:
            self.title += data
        elif not self.skip:
            self.buf.append(data)

    def close(self):
        super().close()
        self._flush()


def html_blocks(html: str) -> _HtmlBlocks:
    parser = _HtmlBlocks()
    parser.feed(html)
    parser.close()
    return parser


def parse_html(html: str, filename: str) -> Parsed:
    parser = html_blocks(html)
    main_chars = sum(len(t) for _, t in parser.main_blocks)
    blocks = parser.main_blocks if main_chars >= 200 else parser.all
    title = re.sub(r"\s+", " ", parser.title).strip()
    if not title:
        title = next((t for lvl, t in blocks if lvl == 1), "") or _stem(filename)
    items: List[Tuple[Optional[int], str]] = []
    for level, text in blocks:
        items.append((level, text))
        if level is None:
            items.append((None, ""))
    props = _props(**parser.meta)
    doc_date = next((props[k][:10] for k in ("modified", "published") if isinstance(props.get(k), str)
                     and re.match(r"\d{4}-\d{2}-\d{2}", props[k])), None)
    return Parsed(title=title, units=_units_from_headed_lines(items), doc_type="html", doc_date=doc_date,
                  properties=props)


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------

_PDF_HEADING = re.compile(r"^(\d{1,2}(?:\.\d{1,2}){0,2})\.?\s+([A-ZА-ЯЁ][^\n]{2,90})$")


def _pdf_date(value: Any) -> Optional[str]:
    if isinstance(value, datetime):
        return value.date().isoformat()
    return None


def parse_pdf(data: bytes, filename: str) -> Parsed:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ExtractError("для PDF нужен пакет pypdf") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:  # noqa: BLE001
                raise ExtractError("PDF зашифрован паролем") from exc
        pages = [(page.extract_text() or "") for page in reader.pages]
    except ExtractError:
        raise
    except Exception as exc:  # noqa: BLE001 — pypdf бросает разные исключения
        raise ExtractError(f"не удалось прочитать PDF: {exc}") from exc
    if sum(len(p.strip()) for p in pages) < 20:
        raise ExtractError("нет текстового слоя — нужно распознавание текста (OCR)")

    # Колонтитулы: строки, повторяющиеся в начале/конце большинства страниц.
    def key(line: str) -> str:
        return re.sub(r"\d+", "#", line.strip().lower())

    edge = Counter()
    split_pages = [[ln for ln in p.replace("\r", "\n").split("\n")] for p in pages]
    for lines in split_pages:
        nonempty = [ln for ln in lines if ln.strip()]
        for ln in set(key(x) for x in nonempty[:2] + nonempty[-2:]):
            edge[ln] += 1
    repeated = {k for k, n in edge.items() if len(pages) >= 3 and n >= max(2, len(pages) * 0.5) and k}

    items: List[Tuple[Optional[int], str, int]] = []
    for page_no, lines in enumerate(split_pages, start=1):
        nonempty_idx = [i for i, ln in enumerate(lines) if ln.strip()]
        edges = set(nonempty_idx[:2] + nonempty_idx[-2:])
        kept = [ln for i, ln in enumerate(lines) if not (i in edges and key(ln) in repeated)]
        text = "\n".join(kept)
        text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
        text = re.sub(r"(?<![.!?:;])\n(?=[a-zа-яё(«\"])", " ", text)
        for line in text.split("\n"):
            stripped = line.strip()
            match = _PDF_HEADING.match(stripped)
            if match and not stripped.endswith((".", ",", ";")):
                level = min(3, match.group(1).count(".") + 1)
                items.append((level, stripped, page_no))
            else:
                items.append((None, line, page_no))

    units: List[Unit] = []
    path = _SectionPath()
    section, buf, buf_page = "", [], 1
    title = None

    def flush() -> None:
        text = _clean("\n".join(buf))
        if text:
            units.append(Unit(text=text, section=section, page=buf_page))

    for level, line, page_no in items:
        if level is not None:
            flush()
            buf, buf_page = [line], page_no
            section = path.set(level, line)
            title = title or (line if level == 1 else None)
        elif page_no != buf_page:
            flush()
            buf, buf_page = [line], page_no
        else:
            buf.append(line)
    flush()
    meta = reader.metadata or {}
    meta_title = (getattr(meta, "title", None) or "").strip() if meta else ""
    doc_date = None
    if meta:
        doc_date = _pdf_date(getattr(meta, "modification_date", None)) or _pdf_date(getattr(meta, "creation_date", None))
    properties = _props(
        pages=len(pages),
        author=getattr(meta, "author", None) if meta else None,
        subject=getattr(meta, "subject", None) if meta else None,
        keywords=meta.get("/Keywords") if meta else None,
        creator=getattr(meta, "creator", None) if meta else None,
        producer=getattr(meta, "producer", None) if meta else None,
        created=_pdf_date(getattr(meta, "creation_date", None)) if meta else None,
        modified=_pdf_date(getattr(meta, "modification_date", None)) if meta else None,
    )
    return Parsed(title=meta_title or title or _stem(filename), units=units, doc_type="pdf", doc_date=doc_date,
                  properties=properties)


# ---------------------------------------------------------------------------
# Офисные форматы
# ---------------------------------------------------------------------------

_DOCX_HEADING = re.compile(r"^(?:heading|заголовок)\s*(\d)$", re.I)


def parse_docx(data: bytes, filename: str) -> Parsed:
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover
        raise ExtractError("для DOCX нужен пакет python-docx") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractError(f"не удалось прочитать DOCX: {exc}") from exc
    items: List[Tuple[Optional[int], str]] = []
    title = None
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                items.append((None, ""))
                continue
            style = (para.style.name if para.style is not None else "") or ""
            match = _DOCX_HEADING.match(style.strip())
            if style.lower() == "title":
                title = title or text
                items.append((1, text))
            elif match:
                level = int(match.group(1))
                title = title or (text if level == 1 else None)
                items.append((level, text))
            elif "list" in style.lower() or "спис" in style.lower():
                items.append((None, "- " + text))
            else:
                items.append((None, text))
                items.append((None, ""))
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells: List[str] = []
                for cell in row.cells:
                    value = cell.text.strip().replace("\n", " ")
                    if not cells or cells[-1] != value:
                        cells.append(value)
                if any(cells):
                    items.append((None, " | ".join(cells)))
            items.append((None, ""))
    props = document.core_properties
    doc_date = props.modified.date().isoformat() if props.modified else None
    properties = _props(
        author=props.author, last_modified_by=props.last_modified_by, subject=props.subject,
        keywords=props.keywords, category=props.category, comments=props.comments,
        created=props.created, modified=props.modified,
        paragraphs=sum(1 for level, text in items if level is None and text),
        tables=len(document.tables),
    )
    return Parsed(title=(props.title or "").strip() or title or _stem(filename),
                  units=_units_from_headed_lines(items), doc_type="docx", doc_date=doc_date,
                  doc_version=(props.version or "").strip() or None, properties=properties)


def parse_pptx(data: bytes, filename: str) -> Parsed:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover
        raise ExtractError("для PPTX нужен пакет python-pptx") from exc
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ExtractError(f"не удалось прочитать PPTX: {exc}") from exc
    units: List[Unit] = []
    first_title = None
    for number, slide in enumerate(prs.slides, start=1):
        title_shape = slide.shapes.title
        slide_title = (title_shape.text.strip() if title_shape is not None and title_shape.has_text_frame else "")
        first_title = first_title or slide_title or None
        parts: List[str] = [slide_title] if slide_title else []
        for shape in slide.shapes:
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                parts.append("Заметки: " + notes)
        text = _clean("\n".join(parts))
        if text:
            section = f"Слайд {number}" + (f": {slide_title}" if slide_title else "")
            units.append(Unit(text=text, section=section, page=number))
    props = prs.core_properties
    doc_date = props.modified.date().isoformat() if props.modified else None
    properties = _props(
        slides=len(prs.slides), author=props.author, last_modified_by=props.last_modified_by,
        subject=props.subject, keywords=props.keywords, category=props.category,
        created=props.created, modified=props.modified,
    )
    return Parsed(title=(props.title or "").strip() or first_title or _stem(filename), units=units,
                  doc_type="pptx", doc_date=doc_date, doc_version=(props.version or "").strip() or None,
                  properties=properties)


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_xlsx(data: bytes, filename: str) -> Parsed:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ExtractError("для XLSX нужен пакет openpyxl") from exc
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExtractError(f"не удалось прочитать XLSX: {exc}") from exc
    units: List[Unit] = []
    total_rows = 0
    for number, ws in enumerate(wb.worksheets, start=1):
        header: Optional[List[str]] = None
        lines: List[str] = []
        for row in ws.iter_rows(values_only=True):
            values = [_cell(v) for v in row]
            if not any(values):
                continue
            if header is None:
                header = [v or f"Столбец {i + 1}" for i, v in enumerate(values)]
                lines.append(" | ".join(v for v in values if v))
                continue
            pairs = [f"{header[i] if i < len(header) else f'Столбец {i + 1}'}: {v}" for i, v in enumerate(values) if v]
            lines.append("; ".join(pairs))
            total_rows += 1
        if lines:
            units.append(Unit(text="\n".join(lines), section=f"Лист «{ws.title}»", page=number))
    sheets = [ws.title for ws in wb.worksheets]
    props = wb.properties
    wb.close()
    properties = _props(
        sheets=len(sheets), sheet_names=sheets, rows=total_rows, author=props.creator,
        last_modified_by=props.lastModifiedBy, subject=props.subject, keywords=props.keywords,
        category=props.category, created=props.created, modified=props.modified,
    )
    doc_date = properties.get("modified") or properties.get("created")
    return Parsed(title=(props.title or "").strip() or _stem(filename), units=units, doc_type="xlsx",
                  doc_date=doc_date, properties=properties)


# ---------------------------------------------------------------------------
# JSONL
# ---------------------------------------------------------------------------


def _get_path(record: Any, dotted: str) -> Any:
    value = record
    for part in dotted.split("."):
        if isinstance(value, dict) and part in value:
            value = value[part]
        else:
            return None
    return value


def _scalar(value: Any) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return "" if value is None else str(value)


def parse_jsonl(text: str, filename: str, text_fields: List[str], metadata_fields: List[str]) -> Parsed:
    units: List[Unit] = []
    bad = 0
    for number, line in enumerate(text.splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        try:
            record = json.loads(line)
        except ValueError:
            bad += 1
            continue
        if not isinstance(record, dict):
            record = {"value": record}
        meta = {f: _get_path(record, f) for f in metadata_fields if _get_path(record, f) is not None}
        if "id" in record and "id" not in meta and not isinstance(record["id"], (dict, list)):
            meta["id"] = record["id"]
        fields = text_fields or [k for k, v in record.items()
                                 if k not in metadata_fields and k != "id" and isinstance(v, (str, int, float))]
        parts = []
        for f in fields:
            value = _scalar(_get_path(record, f)).strip()
            if value:
                parts.append(value if len(fields) == 1 else f"{f}: {value}")
        if not parts:
            continue
        section = f"Запись {meta['id']}" if "id" in meta else f"Запись {number}"
        units.append(Unit(text="\n".join(parts), section=section, atomic=True,
                          metadata={k: v for k, v in meta.items() if isinstance(v, (str, int, float, bool))}))
    parsed = Parsed(title=_stem(filename), units=units, doc_type="jsonl",
                    properties=_props(records=len(units), text_fields=text_fields or None,
                                      metadata_fields=metadata_fields or None))
    if bad:
        parsed.warnings.append(f"пропущено строк с некорректным JSON: {bad}")
    return parsed


# ---------------------------------------------------------------------------
# Точка входа
# ---------------------------------------------------------------------------


def parse_document(data: bytes, filename: str, file_types: FileTypes, charset: Optional[str] = None) -> Parsed:
    kind = resolve_kind(filename, file_types)
    if kind is None:
        raise ExtractError(f"формат {extension_of(filename) or 'без расширения'} не поддерживается")
    if kind == "pdf":
        parsed = parse_pdf(data, filename)
    elif kind == "docx":
        parsed = parse_docx(data, filename)
    elif kind == "pptx":
        parsed = parse_pptx(data, filename)
    elif kind == "xlsx":
        parsed = parse_xlsx(data, filename)
    else:
        text = decode_text(data, charset)
        if kind == "markdown":
            parsed = parse_markdown(text, filename)
        elif kind == "rst":
            parsed = parse_rst(text, filename)
        elif kind == "html":
            parsed = parse_html(text, filename)
        elif kind == "jsonl":
            parsed = parse_jsonl(text, filename, file_types.jsonl_text_fields, file_types.jsonl_metadata_fields)
        elif kind == "code":
            language, units = code_structure.split_code(text, filename)
            declarations = [u.section for u in units if u.section]
            parsed = Parsed(title=PurePosixPath(filename.replace("\\", "/")).name, units=units,
                            doc_type="code", code_language=language,
                            properties=_props(language=language, lines=text.count("\n") + 1,
                                              declarations=len(declarations) or None))
        else:
            parsed = parse_text(text, filename)
    parsed.units = [u for u in parsed.units if u.text.strip()]
    parsed.properties.setdefault("sections", len({u.section for u in parsed.units if u.section}) or None)
    parsed.properties = {k: v for k, v in parsed.properties.items() if v is not None}
    if not parsed.units:
        raise ExtractError("в документе нет текста")
    return parsed
